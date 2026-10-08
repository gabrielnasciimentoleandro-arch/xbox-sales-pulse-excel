#!/usr/bin/env python3
"""Valida dados, fórmulas, gráficos e integridade do Xbox Sales Pulse."""

from __future__ import annotations

from csv import reader
from hashlib import sha256
from math import isclose
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / "xbox_sales_pulse.xlsx"
SOURCE = ROOT / "data" / "base.xlsx"
CSV_FILE = ROOT / "data" / "assinaturas_xbox.csv"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"
SOURCE_SHA256 = "0b60de55f2e54ef2759b6363fe1ef1436ea7ce3de5493ddd5f0e987ac8052a13"
EXPECTED_SHEETS = ["Dashboard", "Análise", "Base", "Dicionário", "Guia", "Apoio"]
EXPECTED_NAMES = {
    "lista_planos", "lista_tipos", "lista_renovacao",
    "filtro_plano", "filtro_tipo", "filtro_renovacao",
    "base_id", "base_plano", "base_inicio", "base_renovacao",
    "base_preco", "base_tipo", "base_ea", "base_minecraft",
    "base_cupom", "base_total",
}


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def main() -> None:
    assert SOURCE.exists(), "Base original ausente"
    assert FILE.exists(), "Dashboard ausente"
    assert CSV_FILE.exists(), "CSV normalizado ausente"
    assert file_hash(SOURCE) == SOURCE_SHA256, "A base original foi alterada"

    with ZipFile(FILE) as zf:
        assert zf.testzip() is None, "Entrada corrompida no XLSX"
        names = set(zf.namelist())
        assert "xl/workbook.xml" in names
        assert "docProps/core.xml" in names
        assert len([name for name in names if name.startswith("xl/charts/chart")]) == 4
        assert not any(name.startswith("xl/externalLinks/") for name in names)
        assert not any(name.endswith("vbaProject.bin") for name in names)

    wb = load_workbook(FILE, data_only=False)
    cached = load_workbook(FILE, data_only=True)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb["Apoio"].sheet_state == "hidden"
    assert wb.properties.creator == AUTHOR
    assert set(wb.defined_names) == EXPECTED_NAMES

    assert len(wb["Dashboard"]._charts) == 4
    assert len(wb["Dashboard"].data_validations.dataValidation) == 3
    assert len(wb["Base"].tables) == 1
    assert "tblAssinaturas" in wb["Base"].tables
    assert wb["Base"].tables["tblAssinaturas"].ref == "A4:P299"

    # Base normalizada.
    base = wb["Base"]
    ids = [base.cell(row, 1).value for row in range(5, 300)]
    assert len(ids) == 295
    assert len(set(ids)) == 295
    assert min(ids) == 3231
    assert max(ids) == 3525
    assert base["N5"].value == "=F5+I5+K5-L5"
    assert base["O5"].value == '=IF(ABS(M5-N5)<0.001,"OK","REVIEW")'
    assert base["P5"].value == "=DATE(YEAR(D5),MONTH(D5),1)"

    source_wb = load_workbook(SOURCE, data_only=True, read_only=True)
    source_ws = next(ws for ws in source_wb.worksheets if ws["A1"].value == "Subscriber ID")
    source_rows = source_ws.iter_rows(min_row=2, max_row=296, max_col=13, values_only=True)
    monetary_columns = {6, 9, 11, 12, 13}
    for source_row_number, (source_values, output_row) in enumerate(
        zip(source_rows, range(5, 300)), start=2
    ):
        for col, source_value in enumerate(source_values, start=1):
            expected_value = 0 if col in monetary_columns and source_value in (None, "-") else source_value
            assert expected_value == base.cell(output_row, col).value, (
                f"Divergência da fonte em linha {source_row_number}, coluna {col}"
            )

    assert sum(float(base.cell(row, 13).value) for row in range(5, 300)) == 7633
    assert sum(float(base.cell(row, 9).value) for row in range(5, 300)) == 2940
    assert sum(float(base.cell(row, 11).value) for row in range(5, 300)) == 3880
    assert sum(float(base.cell(row, 12).value) for row in range(5, 300)) == 2122

    cbase = cached["Base"]
    assert all(cbase.cell(row, 15).value == "OK" for row in range(5, 300))
    assert all(isclose(cbase.cell(row, 14).value, cbase.cell(row, 13).value, abs_tol=1e-9)
               for row in range(5, 300))

    # Fórmulas dinâmicas e perguntas de negócio.
    analysis = wb["Análise"]
    assert analysis["B6"].value.startswith("=SUMIFS(")
    assert "filtro_plano" in analysis["B6"].value
    assert "filtro_tipo" in analysis["B6"].value
    assert "filtro_renovacao" in analysis["B6"].value
    assert analysis["B9"].value.startswith("=IFERROR(COUNTIFS(")
    assert "Annual" in analysis["B17"].value
    assert "Yes" in analysis["B18"].value
    assert "No" in analysis["B19"].value

    can = cached["Análise"]
    expected = {
        "B6": 7633.0,
        "B7": 295,
        "B8": 7633 / 295,
        "B9": 148 / 295,
        "B10": 2935.0,
        "B11": 2940.0,
        "B12": 3880.0,
        "B13": 2122.0,
        "B14": 1.0,
        "B17": 1754.0,
        "B18": 1537.0,
        "B19": 217.0,
        "B20": 2940.0,
        "B21": 3880.0,
    }
    for coordinate, value in expected.items():
        assert isclose(float(can[coordinate].value), value, abs_tol=1e-9), coordinate

    # Distribuições conhecidas da base.
    assert [can[f"F{row}"].value for row in range(26, 29)] == [444, 1801, 5388]
    assert [can[f"G{row}"].value for row in range(26, 29)] == [101, 96, 98]
    assert [can[f"F{row}"].value for row in range(33, 36)] == [3571, 2308, 1754]
    assert [can[f"G{row}"].value for row in range(33, 36)] == [139, 85, 71]
    assert [can[f"J{row}"].value for row in range(26, 28)] == [3786, 3847]
    assert [can[f"K{row}"].value for row in range(26, 28)] == [148, 147]

    cdash = cached["Dashboard"]
    assert cdash["A11"].value == 7633
    assert cdash["D11"].value == 295
    assert isclose(cdash["G11"].value, 7633 / 295, abs_tol=1e-9)
    assert isclose(cdash["J11"].value, 148 / 295, abs_tol=1e-9)
    assert cdash["M11"].value == 1
    assert wb["Dashboard"]["C6"].value == "Todos"
    assert wb["Dashboard"]["G6"].value == "Todos"
    assert wb["Dashboard"]["K6"].value == "Todos"

    # Nenhuma fórmula com erro de construção.
    formula_count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    assert "=IFERROR(=" not in cell.value
                    assert "(=" not in cell.value
                    assert "[" not in cell.value, f"Possível vínculo externo: {ws.title}!{cell.coordinate}"
    assert formula_count >= 950

    excel_errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    for ws in cached.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.value not in excel_errors, f"Erro em {ws.title}!{cell.coordinate}"

    internal_links = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink and cell.hyperlink.target.startswith("#"):
                    internal_links += 1
    assert internal_links >= 25

    with CSV_FILE.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(reader(handle))
    assert len(rows) == 296
    assert rows[0][0] == "Subscriber ID"
    assert rows[-1][0] == "3525"

    for image_name in ("dashboard-geral.png", "data-view.png"):
        path = ROOT / "docs" / image_name
        assert path.exists(), f"Prévia ausente: {image_name}"
        with Image.open(path) as image:
            assert image.width >= 2500
            assert image.height >= 1400

    presentation = ROOT / "docs" / "xbox-sales-pulse-apresentacao.pdf"
    assert presentation.exists(), "Apresentação PDF ausente"
    pdf_bytes = presentation.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-"), "Cabeçalho PDF inválido"
    assert pdf_bytes.rstrip().endswith(b"%%EOF"), "PDF incompleto"
    assert pdf_bytes.count(b"/Type /Page") == 5, "A apresentação deve ter quatro páginas"

    print("VALIDAÇÃO APROVADA")
    print(f"Arquivo: {FILE.name}")
    print(f"Autoria: {wb.properties.creator}")
    print(f"Abas: {len(wb.sheetnames)} ({', '.join(wb.sheetnames)})")
    print("Base: 295 registros, 295 IDs únicos e 295 valores conferidos")
    print("Receita: $ 7.633,00 | Ticket médio: $ 25,87 | Auto renovação: 50,2%")
    print("Planos anuais: $ 1.754,00 ($ 1.537,00 com renovação; $ 217,00 sem)")
    print("EA Play: $ 2.940,00 | Minecraft: $ 3.880,00 | Cupons: $ 2.122,00")
    print(f"Gráficos: 4 | Filtros: 3 | Fórmulas: {formula_count} | Links internos: {internal_links}")
    print(f"SHA-256 base original: {SOURCE_SHA256}")
    print("Integridade ZIP/XML, caches, nomes, tabela, CSV e imagens: OK")
    print("Macros e vínculos externos de pasta de trabalho: AUSENTES")


if __name__ == "__main__":
    main()
