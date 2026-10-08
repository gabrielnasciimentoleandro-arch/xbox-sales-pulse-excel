#!/usr/bin/env python3
"""Gera o dashboard autoral Xbox Sales Pulse em Excel.

A base didática é a base.xlsx fornecida no desafio da DIO. O projeto cria
uma visualização própria, sem reutilizar o dashboard final de referência.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import csv
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "base.xlsx"
OUTPUT = ROOT / "xbox_sales_pulse.xlsx"
CSV_OUTPUT = ROOT / "data" / "assinaturas_xbox.csv"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"
AUTHOR_HANDLE = "gabrielnasciimentoleandro-arch"

# Paleta autoral inspirada no universo gamer, sem reutilizar logos oficiais.
INK = "0B0F14"
DARK = "111827"
SLATE = "1F2937"
GREEN = "22C55E"
LIME = "9BC848"
MINT = "2AE6B1"
BLUE = "38BDF8"
PURPLE = "8B5CF6"
AMBER = "F59E0B"
RED = "EF4444"
WHITE = "FFFFFF"
TEXT = "172033"
MUTED = "64748B"
LIGHT = "F4F7FA"
PALE_GREEN = "DCFCE7"
PALE_RED = "FEE2E2"
BORDER = "CBD5E1"

thin = Side(style="thin", color=BORDER)
medium_green = Side(style="medium", color=GREEN)
CURRENCY = '$ #,##0.00;[Red]-$ #,##0.00'
PERCENT = '0.0%'
INTEGER = '0'
DATE_FMT = 'dd/mm/yyyy'
MONTH_FMT = 'mmm/yyyy'

HEADERS = [
    "Subscriber ID",
    "Name",
    "Plan",
    "Start Date",
    "Auto Renewal",
    "Subscription Price",
    "Subscription Type",
    "EA Play Season Pass",
    "EA Play Season Pass Price",
    "Minecraft Season Pass",
    "Minecraft Season Pass Price",
    "Coupon Value",
    "Total Value",
]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def border(accent: bool = False) -> Border:
    return Border(
        left=medium_green if accent else thin,
        right=thin,
        top=thin,
        bottom=thin,
    )


def add_name(wb: Workbook, name: str, reference: str) -> None:
    wb.defined_names.add(DefinedName(name=name, attr_text=reference))


def merge_box(ws, cell_range: str, value, *, bg: str = WHITE, color: str = TEXT,
              size: int = 10, bold: bool = False, align: str = "center",
              number_format: str | None = None, hyperlink: str | None = None) -> None:
    ws.merge_cells(cell_range)
    anchor = cell_range.split(":")[0]
    cell = ws[anchor]
    cell.value = value
    cell.fill = fill(bg)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = border()
    if number_format:
        cell.number_format = number_format
    if hyperlink:
        cell.hyperlink = hyperlink


def find_source_sheet(workbook) -> object:
    for ws in workbook.worksheets:
        values = [ws.cell(1, col).value for col in range(1, 14)]
        if values and values[0] == "Subscriber ID" and values[2] == "Plan":
            return ws
    raise ValueError("A aba de dados da base.xlsx não foi encontrada.")


def load_records() -> list[list]:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Base não encontrada: {SOURCE}")
    source_wb = load_workbook(SOURCE, data_only=True, read_only=True)
    ws = find_source_sheet(source_wb)
    headers = [ws.cell(1, col).value for col in range(1, 14)]
    assert headers == [
        "Subscriber ID", "Name", "Plan", "Start Date", "Auto Renewal",
        "Subscription Price", "Subscription Type", "EA Play Season Pass",
        "EA Play Season Pass\nPrice", "Minecraft Season Pass",
        "Minecraft Season Pass Price", "Coupon Value", "Total Value",
    ]

    records: list[list] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if values[0] is None:
            continue
        row = list(values[:13])
        if row[8] in (None, "-"):
            row[8] = 0
        for index in (5, 8, 10, 11, 12):
            row[index] = float(row[index] or 0)
        records.append(row)
    assert len(records) == 295
    return records


def write_normalized_csv(records: list[list]) -> None:
    CSV_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with CSV_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for row in records:
            out = list(row)
            if isinstance(out[3], (date, datetime)):
                out[3] = out[3].strftime("%Y-%m-%d")
            writer.writerow(out)


def summarize(records: list[list]) -> dict:
    summary: dict = {}
    summary["count"] = len(records)
    summary["start"] = min(row[3] for row in records)
    summary["end"] = max(row[3] for row in records)
    summary["revenue"] = sum(row[12] for row in records)
    summary["base"] = sum(row[5] for row in records)
    summary["ea"] = sum(row[8] for row in records)
    summary["minecraft"] = sum(row[10] for row in records)
    summary["coupons"] = sum(row[11] for row in records)
    summary["ticket"] = summary["revenue"] / summary["count"]
    summary["auto_yes"] = sum(1 for row in records if row[4] == "Yes")
    summary["auto_rate"] = summary["auto_yes"] / summary["count"]

    summary["plans"] = {}
    for plan in ("Core", "Standard", "Ultimate"):
        subset = [row for row in records if row[2] == plan]
        summary["plans"][plan] = {
            "count": len(subset),
            "revenue": sum(row[12] for row in subset),
            "ea": sum(row[8] for row in subset),
            "minecraft": sum(row[10] for row in subset),
        }

    summary["types"] = {}
    for kind in ("Monthly", "Quarterly", "Annual"):
        subset = [row for row in records if row[6] == kind]
        summary["types"][kind] = {
            "count": len(subset),
            "revenue": sum(row[12] for row in subset),
        }

    summary["renewal"] = {}
    for status in ("Yes", "No"):
        subset = [row for row in records if row[4] == status]
        summary["renewal"][status] = {
            "count": len(subset),
            "revenue": sum(row[12] for row in subset),
        }

    summary["annual_total"] = summary["types"]["Annual"]["revenue"]
    summary["annual_yes"] = sum(row[12] for row in records if row[6] == "Annual" and row[4] == "Yes")
    summary["annual_no"] = sum(row[12] for row in records if row[6] == "Annual" and row[4] == "No")

    monthly: dict[int, dict[str, float | int]] = {}
    for month in range(1, 13):
        subset = [row for row in records if row[3].month == month]
        monthly[month] = {
            "count": len(subset),
            "revenue": sum(row[12] for row in subset),
        }
    summary["monthly"] = monthly
    summary["validation_ok"] = sum(
        1 for row in records
        if abs((row[5] + row[8] + row[10] - row[11]) - row[12]) < 1e-9
    )
    return summary


def setup_sheet(ws, title: str, subtitle: str, *, freeze: str | None = None,
                zoom: int = 90) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.freeze_panes = freeze
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = .25
    ws.page_margins.right = .25
    ws.page_margins.top = .35
    ws.page_margins.bottom = .35
    ws.oddFooter.center.text = f"Xbox Sales Pulse • {AUTHOR}"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = MUTED

    for col in range(1, 15):
        ws.column_dimensions[chr(64 + col)].width = 13
    for row in range(1, 4):
        ws.row_dimensions[row].height = 24
        for col in range(1, 15):
            ws.cell(row, col).fill = fill(INK)

    ws.merge_cells("A1:I2")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=23, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells("J1:N1")
    ws["J1"] = "PROJETO AUTORAL • DIO"
    ws["J1"].font = Font(name="Aptos", size=9, bold=True, color=GREEN)
    ws["J1"].alignment = Alignment(horizontal="right", vertical="center")

    navigation = [
        ("J2", "DASHBOARD", "Dashboard"),
        ("K2", "ANÁLISE", "Análise"),
        ("L2", "BASE", "Base"),
        ("M2", "DICIONÁRIO", "Dicionário"),
        ("N2", "GUIA", "Guia"),
    ]
    for coordinate, label, target in navigation:
        cell = ws[coordinate]
        cell.value = label
        cell.fill = fill(GREEN if ws.title == target else SLATE)
        cell.font = Font(name="Aptos", size=7, bold=True, color=INK if ws.title == target else WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=medium_green)
        cell.hyperlink = f"#'{target}'!A1"

    ws.merge_cells("A3:N3")
    ws["A3"] = subtitle
    ws["A3"].fill = fill(LIGHT)
    ws["A3"].font = Font(name="Aptos", size=9, italic=True, color=MUTED)
    ws["A3"].alignment = Alignment(vertical="center")
    ws["A3"].border = Border(bottom=thin)


def style_input(cell, number_format: str | None = None) -> None:
    cell.fill = fill(PALE_GREEN)
    cell.font = Font(name="Aptos", size=11, bold=True, color=TEXT)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border(accent=True)
    cell.protection = Protection(locked=False)
    if number_format:
        cell.number_format = number_format


def style_formula(cell, number_format: str | None = None, align: str = "right") -> None:
    cell.fill = fill(WHITE)
    cell.font = Font(name="Aptos", size=10, bold=True, color=TEXT)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = border()
    if number_format:
        cell.number_format = number_format


def add_dropdown(ws, target: str, named_range: str, prompt: str) -> None:
    dv = DataValidation(type="list", formula1=f"={named_range}", allow_blank=False)
    dv.errorTitle = "Filtro inválido"
    dv.error = "Escolha uma opção disponível na lista."
    dv.promptTitle = "Filtro do dashboard"
    dv.prompt = prompt
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(target)


def sumifs(sum_range: str, extra: list[tuple[str, str]] | None = None,
           include_filters: bool = True) -> str:
    pairs: list[tuple[str, str]] = []
    if include_filters:
        pairs.extend([
            ("Base!$C$5:$C$299", 'IF(filtro_plano="Todos","*",filtro_plano)'),
            ("Base!$G$5:$G$299", 'IF(filtro_tipo="Todos","*",filtro_tipo)'),
            ("Base!$E$5:$E$299", 'IF(filtro_renovacao="Todos","*",filtro_renovacao)'),
        ])
    if extra:
        pairs.extend(extra)
    args = [sum_range]
    for range_ref, criterion in pairs:
        args.extend([range_ref, criterion])
    return "=SUMIFS(" + ",".join(args) + ")"


def countifs(extra: list[tuple[str, str]] | None = None,
             include_filters: bool = True) -> str:
    pairs: list[tuple[str, str]] = []
    if include_filters:
        pairs.extend([
            ("Base!$C$5:$C$299", 'IF(filtro_plano="Todos","*",filtro_plano)'),
            ("Base!$G$5:$G$299", 'IF(filtro_tipo="Todos","*",filtro_tipo)'),
            ("Base!$E$5:$E$299", 'IF(filtro_renovacao="Todos","*",filtro_renovacao)'),
        ])
    if extra:
        pairs.extend(extra)
    args: list[str] = []
    for range_ref, criterion in pairs:
        args.extend([range_ref, criterion])
    return "=COUNTIFS(" + ",".join(args) + ")"


def card(ws, start_col: int, end_col: int, label_row: int, value_start: int,
         value_end: int, label: str, formula: str, fmt: str, accent: str) -> str:
    start = chr(64 + start_col)
    end = chr(64 + end_col)
    merge_box(ws, f"{start}{label_row}:{end}{label_row}", label,
              bg=SLATE, color=WHITE, size=8, bold=True)
    merge_box(ws, f"{start}{value_start}:{end}{value_end}", formula,
              bg=WHITE, color=accent, size=18, bold=True, number_format=fmt)
    return f"{start}{value_start}"


def style_chart(chart, title: str, width: float = 13.2, height: float = 7.2,
                data_labels: bool = False) -> None:
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height
    chart.legend.position = "b"
    if data_labels:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True


def build_dashboard(wb: Workbook, summary: dict) -> None:
    ws = wb.active
    ws.title = "Dashboard"
    setup_sheet(
        ws,
        "XBOX SALES PULSE",
        "Dashboard autoral de assinaturas: receita, planos, recorrência e adicionais em uma visão executiva.",
        freeze="A4",
        zoom=75,
    )
    for col in range(1, 15):
        ws.column_dimensions[chr(64 + col)].width = 13.5
    for row in range(4, 59):
        ws.row_dimensions[row].height = 21

    merge_box(ws, "A4:N4",
              '="Base didática: "&Análise!B7&" assinaturas • período "&TEXT(MIN(Base!D5:D299),"dd/mm/yyyy")&" a "&TEXT(MAX(Base!D5:D299),"dd/mm/yyyy")&" • valores em US$"',
              bg=DARK, color=WHITE, size=9, bold=True)

    merge_box(ws, "A5:N5", "FILTROS INTERATIVOS — selecione uma opção nas células verdes",
              bg=SLATE, color=WHITE, size=9, bold=True)
    merge_box(ws, "A6:B7", "PLANO", bg=LIGHT, color=MUTED, size=9, bold=True)
    merge_box(ws, "C6:D7", "Todos", bg=PALE_GREEN, color=TEXT, size=11, bold=True)
    style_input(ws["C6"])
    merge_box(ws, "E6:F7", "PERIODICIDADE", bg=LIGHT, color=MUTED, size=9, bold=True)
    merge_box(ws, "G6:H7", "Todos", bg=PALE_GREEN, color=TEXT, size=11, bold=True)
    style_input(ws["G6"])
    merge_box(ws, "I6:J7", "RENOVAÇÃO", bg=LIGHT, color=MUTED, size=9, bold=True)
    merge_box(ws, "K6:L7", "Todos", bg=PALE_GREEN, color=TEXT, size=11, bold=True)
    style_input(ws["K6"])
    merge_box(ws, "M6:N7", "Para limpar:\nescolha Todos", bg=DARK, color=GREEN, size=8, bold=True)

    add_dropdown(ws, "C6", "lista_planos", "Filtre Core, Standard, Ultimate ou Todos.")
    add_dropdown(ws, "G6", "lista_tipos", "Filtre Monthly, Quarterly, Annual ou Todos.")
    add_dropdown(ws, "K6", "lista_renovacao", "Filtre Yes, No ou Todos.")
    ws["C6"].comment = Comment("Todos os cartões e gráficos respondem aos três filtros.", "Xbox Sales Pulse")

    merge_box(ws, "A8:N8", '="Filtros ativos: Plano = "&filtro_plano&" • Periodicidade = "&filtro_tipo&" • Renovação = "&filtro_renovacao',
              bg=PALE_GREEN, color=TEXT, size=9, bold=True)

    card(ws, 1, 3, 10, 11, 13, "RECEITA TOTAL", "=Análise!B6", CURRENCY, GREEN)
    card(ws, 4, 6, 10, 11, 13, "ASSINANTES", "=Análise!B7", INTEGER, BLUE)
    card(ws, 7, 9, 10, 11, 13, "TICKET MÉDIO", "=Análise!B8", CURRENCY, MINT)
    card(ws, 10, 12, 10, 11, 13, "% AUTO RENOVAÇÃO", "=Análise!B9", PERCENT, LIME)
    card(ws, 13, 14, 10, 11, 13, "BASE VALIDADA", "=Análise!B14", PERCENT, PURPLE)

    card(ws, 1, 3, 15, 16, 18, "RECEITA DOS PLANOS", "=Análise!B10", CURRENCY, GREEN)
    card(ws, 4, 6, 15, 16, 18, "EA PLAY", "=Análise!B11", CURRENCY, PURPLE)
    card(ws, 7, 9, 15, 16, 18, "MINECRAFT", "=Análise!B12", CURRENCY, MINT)
    card(ws, 10, 12, 15, 16, 18, "CUPONS CONCEDIDOS", "=Análise!B13", CURRENCY, AMBER)
    card(ws, 13, 14, 15, 16, 18, "REGISTROS", "=Análise!B7", INTEGER, BLUE)

    merge_box(ws, "A20:N20", "VISUALIZAÇÕES DINÂMICAS", bg=INK, color=WHITE, size=9, bold=True)

    monthly = LineChart()
    monthly_data = Reference(wb["Análise"], min_col=2, min_row=25, max_row=37)
    monthly_categories = Reference(wb["Análise"], min_col=1, min_row=26, max_row=37)
    monthly.add_data(monthly_data, titles_from_data=True)
    monthly.set_categories(monthly_categories)
    style_chart(monthly, "Receita por mês", width=13.2, height=7.0, data_labels=False)
    monthly.y_axis.numFmt = '$0'
    monthly.x_axis.number_format = "mmm"
    monthly.series[0].graphicalProperties.line.solidFill = GREEN
    monthly.series[0].graphicalProperties.line.width = 26000
    monthly.series[0].marker.symbol = "circle"
    monthly.series[0].marker.size = 6
    ws.add_chart(monthly, "A21")

    plan_chart = BarChart()
    plan_chart.type = "bar"
    plan_data = Reference(wb["Análise"], min_col=6, min_row=25, max_row=28)
    plan_categories = Reference(wb["Análise"], min_col=5, min_row=26, max_row=28)
    plan_chart.add_data(plan_data, titles_from_data=True)
    plan_chart.set_categories(plan_categories)
    style_chart(plan_chart, "Receita por plano", width=12.2, height=7.0, data_labels=True)
    plan_chart.x_axis.numFmt = '$0'
    plan_chart.series[0].graphicalProperties.solidFill = GREEN
    ws.add_chart(plan_chart, "H21")

    type_chart = BarChart()
    type_chart.type = "col"
    type_data = Reference(wb["Análise"], min_col=6, min_row=32, max_row=35)
    type_categories = Reference(wb["Análise"], min_col=5, min_row=33, max_row=35)
    type_chart.add_data(type_data, titles_from_data=True)
    type_chart.set_categories(type_categories)
    style_chart(type_chart, "Receita por periodicidade", width=12.8, height=7.0, data_labels=True)
    type_chart.y_axis.numFmt = '$0'
    type_chart.series[0].graphicalProperties.solidFill = BLUE
    ws.add_chart(type_chart, "A36")

    addon_chart = BarChart()
    addon_chart.type = "col"
    addon_data = Reference(wb["Análise"], min_col=10, max_col=11, min_row=32, max_row=35)
    addon_categories = Reference(wb["Análise"], min_col=9, min_row=33, max_row=35)
    addon_chart.add_data(addon_data, titles_from_data=True)
    addon_chart.set_categories(addon_categories)
    style_chart(addon_chart, "Adicionais por plano", width=12.8, height=7.0, data_labels=False)
    addon_chart.y_axis.numFmt = '$0'
    addon_chart.series[0].graphicalProperties.solidFill = PURPLE
    addon_chart.series[1].graphicalProperties.solidFill = MINT
    ws.add_chart(addon_chart, "H36")

    merge_box(ws, "A51:N51", "PERGUNTAS-CHAVE DO DESAFIO — BASE COMPLETA", bg=SLATE,
              color=WHITE, size=9, bold=True)
    challenge_cards = [
        ("A52:C54", "Planos anuais\n$ 1.754,00", GREEN),
        ("D52:F54", "Anual com renovação\n$ 1.537,00", LIME),
        ("G52:I54", "Anual sem renovação\n$ 217,00", RED),
        ("J52:L54", "EA Play\n$ 2.940,00", PURPLE),
        ("M52:N54", "Minecraft\n$ 3.880,00", MINT),
    ]
    for cell_range, text, accent in challenge_cards:
        merge_box(ws, cell_range, text, bg=WHITE, color=accent, size=10, bold=True)

    merge_box(ws, "A56:N57",
              f"Projeto autoral: {AUTHOR} • Base didática DIO • dashboard não oficial • Xbox, Game Pass, EA Play e Minecraft pertencem aos respectivos titulares.",
              bg=INK, color=WHITE, size=8, bold=True)
    ws.print_area = "A1:N57"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3


def build_analysis(wb: Workbook) -> None:
    ws = wb.create_sheet("Análise")
    setup_sheet(
        ws,
        "ANÁLISE E CÁLCULOS",
        "Camada transparente de fórmulas que alimenta os cartões e gráficos do Dashboard.",
        freeze="A5",
        zoom=85,
    )
    for col, width in {
        "A": 20, "B": 18, "C": 50, "D": 4,
        "E": 20, "F": 18, "G": 15, "H": 4,
        "I": 20, "J": 18, "K": 18, "L": 18,
        "M": 4, "N": 4,
    }.items():
        ws.column_dimensions[col].width = width

    merge_box(ws, "A4:N4", '="Filtros: "&filtro_plano&" | "&filtro_tipo&" | "&filtro_renovacao',
              bg=PALE_GREEN, color=TEXT, size=9, bold=True)

    headers = ["Indicador dinâmico", "Resultado", "Fórmula / interpretação"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(5, col, value)
        cell.fill = fill(SLATE)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border()

    kpis = [
        (6, "Receita total", sumifs("Base!$M$5:$M$299"), CURRENCY, "Soma de Total Value com os filtros ativos"),
        (7, "Assinantes", countifs(), INTEGER, "Quantidade de registros filtrados"),
        (8, "Ticket médio", "=IFERROR(B6/B7,0)", CURRENCY, "Receita / assinantes"),
        (9, "% auto renovação", f'=IFERROR({countifs([("Base!$E$5:$E$299", "\"Yes\"")])[1:]}/B7,0)', PERCENT, "Participação de Auto Renewal = Yes"),
        (10, "Receita dos planos", sumifs("Base!$F$5:$F$299"), CURRENCY, "Subscription Price"),
        (11, "EA Play", sumifs("Base!$I$5:$I$299"), CURRENCY, "Receita do passe adicional"),
        (12, "Minecraft", sumifs("Base!$K$5:$K$299"), CURRENCY, "Receita do passe adicional"),
        (13, "Cupons", sumifs("Base!$L$5:$L$299"), CURRENCY, "Descontos concedidos"),
        (14, "Base validada", '=COUNTIF(Base!$O$5:$O$299,"OK")/COUNTA(Base!$A$5:$A$299)', PERCENT, "Total Value = plano + adicionais − cupom"),
    ]
    for row, label, formula, fmt, explanation in kpis:
        ws.cell(row, 1, label)
        ws.cell(row, 2, formula)
        ws.cell(row, 3, explanation)
        for col in range(1, 4):
            ws.cell(row, col).border = border()
            ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row, col).fill = fill(WHITE if row % 2 == 0 else LIGHT)
            ws.cell(row, col).font = Font(name="Aptos", size=9, bold=col < 3, color=TEXT)
        ws.cell(row, 2).number_format = fmt
        ws.row_dimensions[row].height = 24

    merge_box(ws, "A16:C16", "PERGUNTAS DE NEGÓCIO — BASE COMPLETA, SEM FILTROS", bg=SLATE,
              color=WHITE, size=9, bold=True)
    challenge = [
        (17, "Receita de planos anuais", sumifs("Base!$M$5:$M$299", [("Base!$G$5:$G$299", '"Annual"')], False)),
        (18, "Anual com auto renovação", sumifs("Base!$M$5:$M$299", [("Base!$G$5:$G$299", '"Annual"'), ("Base!$E$5:$E$299", '"Yes"')], False)),
        (19, "Anual sem auto renovação", sumifs("Base!$M$5:$M$299", [("Base!$G$5:$G$299", '"Annual"'), ("Base!$E$5:$E$299", '"No"')], False)),
        (20, "EA Play Season Pass", "=SUM(Base!$I$5:$I$299)"),
        (21, "Minecraft Season Pass", "=SUM(Base!$K$5:$K$299)"),
    ]
    for row, label, formula in challenge:
        ws.cell(row, 1, label)
        ws.cell(row, 2, formula)
        ws.cell(row, 1).fill = fill(LIGHT)
        ws.cell(row, 2).fill = fill(WHITE)
        ws.cell(row, 1).font = Font(name="Aptos", size=9, bold=True, color=TEXT)
        ws.cell(row, 2).font = Font(name="Aptos", size=10, bold=True, color=GREEN)
        ws.cell(row, 2).number_format = CURRENCY
        ws.cell(row, 1).border = border()
        ws.cell(row, 2).border = border()

    # Séries mensais.
    ws["A25"] = "Mês"
    ws["B25"] = "Receita"
    ws["C25"] = "Assinantes"
    for month in range(1, 13):
        row = 25 + month
        ws.cell(row, 1, date(2024, month, 1))
        ws.cell(row, 1).number_format = MONTH_FMT
        extra = [
            ("Base!$D$5:$D$299", f'">="&A{row}'),
            ("Base!$D$5:$D$299", f'"<"&DATE(YEAR(A{row}),MONTH(A{row})+1,1)'),
        ]
        ws.cell(row, 2, sumifs("Base!$M$5:$M$299", extra))
        ws.cell(row, 3, countifs(extra))
        ws.cell(row, 2).number_format = CURRENCY

    ws["E25"] = "Plano"
    ws["F25"] = "Receita"
    ws["G25"] = "Assinantes"
    for row, plan in enumerate(("Core", "Standard", "Ultimate"), 26):
        ws.cell(row, 5, plan)
        extra = [("Base!$C$5:$C$299", f'E{row}')]
        ws.cell(row, 6, sumifs("Base!$M$5:$M$299", extra))
        ws.cell(row, 7, countifs(extra))
        ws.cell(row, 6).number_format = CURRENCY

    ws["E32"] = "Periodicidade"
    ws["F32"] = "Receita"
    ws["G32"] = "Assinantes"
    for row, kind in enumerate(("Monthly", "Quarterly", "Annual"), 33):
        ws.cell(row, 5, kind)
        extra = [("Base!$G$5:$G$299", f'E{row}')]
        ws.cell(row, 6, sumifs("Base!$M$5:$M$299", extra))
        ws.cell(row, 7, countifs(extra))
        ws.cell(row, 6).number_format = CURRENCY

    ws["I25"] = "Auto Renewal"
    ws["J25"] = "Receita"
    ws["K25"] = "Assinantes"
    for row, status in enumerate(("Yes", "No"), 26):
        ws.cell(row, 9, status)
        extra = [("Base!$E$5:$E$299", f'I{row}')]
        ws.cell(row, 10, sumifs("Base!$M$5:$M$299", extra))
        ws.cell(row, 11, countifs(extra))
        ws.cell(row, 10).number_format = CURRENCY

    ws["I32"] = "Plano"
    ws["J32"] = "EA Play"
    ws["K32"] = "Minecraft"
    ws["L32"] = "Receita total"
    for row, plan in enumerate(("Core", "Standard", "Ultimate"), 33):
        ws.cell(row, 9, plan)
        extra = [("Base!$C$5:$C$299", f'I{row}')]
        ws.cell(row, 10, sumifs("Base!$I$5:$I$299", extra))
        ws.cell(row, 11, sumifs("Base!$K$5:$K$299", extra))
        ws.cell(row, 12, sumifs("Base!$M$5:$M$299", extra))
        for col in (10, 11, 12):
            ws.cell(row, col).number_format = CURRENCY

    # Estilo dos blocos analíticos.
    for start_col, end_col, header_rows, data_rows in [
        (1, 3, [25], range(26, 38)),
        (5, 7, [25, 32], list(range(26, 29)) + list(range(33, 36))),
        (9, 12, [25, 32], list(range(26, 28)) + list(range(33, 36))),
    ]:
        for header_row in header_rows:
            for col in range(start_col, end_col + 1):
                c = ws.cell(header_row, col)
                c.fill = fill(SLATE)
                c.font = Font(name="Aptos", size=8, bold=True, color=WHITE)
                c.alignment = Alignment(horizontal="center")
                c.border = border()
        for row in data_rows:
            for col in range(start_col, end_col + 1):
                c = ws.cell(row, col)
                c.fill = fill(WHITE if row % 2 == 0 else LIGHT)
                c.font = Font(name="Aptos", size=8, color=TEXT)
                c.alignment = Alignment(horizontal="center")
                c.border = border()

    merge_box(ws, "A40:L42",
              "Leitura recomendada: os gráficos respondem aos filtros do Dashboard. A seção de perguntas do desafio usa sempre a base completa para preservar respostas comparáveis. Dezembro contém dados somente até 16/12/2024.",
              bg=PALE_GREEN, color=TEXT, size=9, bold=True)
    ws.print_area = "A1:L42"


def build_base(wb: Workbook, records: list[list]) -> None:
    ws = wb.create_sheet("Base")
    setup_sheet(
        ws,
        "BASE DE ASSINATURAS",
        "Cópia normalizada da base didática: o hífen de EA Play Price foi convertido em zero para permitir cálculos.",
        freeze="A5",
        zoom=75,
    )
    headers = HEADERS + ["Calculated Value", "Validation", "Month"]
    widths = [14, 24, 14, 15, 16, 18, 18, 18, 20, 21, 23, 15, 15, 18, 14, 14]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width

    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.fill = fill(INK)
        cell.font = Font(name="Aptos", size=8, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border()
    ws.row_dimensions[4].height = 40

    for row_index, record in enumerate(records, 5):
        for col, value in enumerate(record, 1):
            ws.cell(row_index, col, value)
        ws.cell(row_index, 14, f"=F{row_index}+I{row_index}+K{row_index}-L{row_index}")
        ws.cell(row_index, 15, f'=IF(ABS(M{row_index}-N{row_index})<0.001,"OK","REVIEW")')
        ws.cell(row_index, 16, f"=DATE(YEAR(D{row_index}),MONTH(D{row_index}),1)")
        for col in range(1, 17):
            cell = ws.cell(row_index, col)
            cell.fill = fill(WHITE if row_index % 2 else LIGHT)
            cell.font = Font(name="Aptos", size=8, color=TEXT)
            cell.border = border()
            cell.alignment = Alignment(vertical="center")
        ws.cell(row_index, 4).number_format = DATE_FMT
        for col in (6, 9, 11, 12, 13, 14):
            ws.cell(row_index, col).number_format = CURRENCY
        ws.cell(row_index, 16).number_format = MONTH_FMT
        ws.row_dimensions[row_index].height = 19

    table = Table(displayName="tblAssinaturas", ref=f"A4:P{4 + len(records)}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    ws.conditional_formatting.add(
        f"O5:O{4 + len(records)}",
        FormulaRule(formula=['$O5="OK"'], fill=fill(PALE_GREEN),
                    font=Font(color=GREEN, bold=True)),
    )
    ws.conditional_formatting.add(
        f"O5:O{4 + len(records)}",
        FormulaRule(formula=['$O5="REVIEW"'], fill=fill(PALE_RED),
                    font=Font(color=RED, bold=True)),
    )
    ws.auto_filter.ref = f"A4:P{4 + len(records)}"
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:P{4 + len(records)}"


def build_dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("Dicionário")
    setup_sheet(
        ws,
        "DICIONÁRIO DE DADOS",
        "Definições, tratamento aplicado e uso de cada campo da base didática.",
        freeze="A6",
        zoom=85,
    )
    for col, width in {"A": 4, "B": 28, "C": 20, "D": 58, "E": 42}.items():
        ws.column_dimensions[col].width = width
    headers = ["Campo", "Tipo", "Descrição", "Tratamento / uso"]
    for col, value in enumerate(headers, 2):
        c = ws.cell(5, col, value)
        c.fill = fill(SLATE)
        c.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center")
        c.border = border()

    rows = [
        ("Subscriber ID", "Inteiro", "Identificador único do registro.", "Chave primária; 295 valores únicos."),
        ("Name", "Texto", "Nome didático do assinante.", "Não é usado como chave porque pode se repetir."),
        ("Plan", "Categoria", "Core, Standard ou Ultimate.", "Filtro e agrupamento de receita."),
        ("Start Date", "Data", "Data inicial da assinatura em 2024.", "Agrupamento mensal e definição do período."),
        ("Auto Renewal", "Categoria", "Yes ou No para renovação automática.", "Filtro, taxa de recorrência e análise anual."),
        ("Subscription Price", "Moeda", "Preço-base da assinatura.", "Componente positivo do valor final."),
        ("Subscription Type", "Categoria", "Monthly, Quarterly ou Annual.", "Filtro e comparação de periodicidade."),
        ("EA Play Season Pass", "Categoria", "Adesão ao passe EA Play.", "Contexto do adicional."),
        ("EA Play Season Pass Price", "Moeda", "Valor do EA Play.", "Hífens da fonte convertidos em zero."),
        ("Minecraft Season Pass", "Categoria", "Adesão ao passe Minecraft.", "Contexto do adicional."),
        ("Minecraft Season Pass Price", "Moeda", "Valor do passe Minecraft.", "Componente positivo do valor final."),
        ("Coupon Value", "Moeda", "Desconto aplicado.", "Componente negativo do valor final."),
        ("Total Value", "Moeda", "Valor líquido do registro.", "Plano + EA + Minecraft − cupom."),
        ("Calculated Value", "Fórmula", "Recalcula a composição da receita.", "Usado para auditoria do Total Value."),
        ("Validation", "Fórmula", "Resultado OK ou REVIEW.", "295 de 295 registros validados."),
        ("Month", "Fórmula/Data", "Primeiro dia do mês de Start Date.", "Apoio para análise temporal."),
    ]
    for row_index, values in enumerate(rows, 6):
        for col, value in enumerate(values, 2):
            c = ws.cell(row_index, col, value)
            c.fill = fill(WHITE if row_index % 2 == 0 else LIGHT)
            c.font = Font(name="Aptos", size=9, bold=col == 2, color=TEXT)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = border()
        ws.row_dimensions[row_index].height = 38

    merge_box(ws, "B24:E26",
              "QUALIDADE DA BASE: 295 IDs únicos, nenhuma linha vazia na tabela e 295 valores finais conferidos pela fórmula Subscription Price + EA Play + Minecraft − Coupon. O período vai de 01/01/2024 a 16/12/2024; dezembro é parcial.",
              bg=PALE_GREEN, color=TEXT, size=10, bold=True)
    merge_box(ws, "B28:E30",
              "FONTE: base didática disponibilizada como material de apoio do desafio DIO. O arquivo original permanece em data/base.xlsx; a aba Base contém a versão normalizada usada nos cálculos.",
              bg=LIGHT, color=MUTED, size=9)
    ws.print_area = "A1:E30"


def build_guide(wb: Workbook) -> None:
    ws = wb.create_sheet("Guia")
    setup_sheet(
        ws,
        "GUIA E MÉTODO ABCDE",
        "Como atualizar, interpretar e reproduzir o Xbox Sales Pulse.",
        freeze="A4",
        zoom=90,
    )
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 84
    for col in "DEFGHIJKLMN":
        ws.column_dimensions[col].width = 3

    sections = [
        (5, "A — Assets", "Identidade visual autoral: verde, grafite, azul e roxo. Nenhum logo oficial foi incorporado ao arquivo."),
        (8, "B — Base", "A aba Base possui 295 assinaturas normalizadas. A fonte original permanece em data/base.xlsx para auditoria."),
        (11, "C — Cálculos", "A aba Análise usa SUMIFS, COUNTIFS, IFERROR, COUNTIF, SUM e fórmulas de datas. Os blocos alimentam todos os gráficos."),
        (14, "D — Dashboard", "Escolha Plano, Periodicidade e Renovação nas células verdes. Cartões e quatro gráficos recalculam automaticamente."),
        (17, "E — Entrega", "O repositório inclui Excel final, base original, CSV normalizado, README, prévias, metodologia e scripts de validação."),
        (20, "Atualizar dados", "Substitua ou acrescente registros na aba Base mantendo as colunas. Se ultrapassar 295 linhas, amplie os intervalos das fórmulas e a tabela."),
        (23, "Perguntas do desafio", "Análise registra a receita anual total, anual com/sem renovação, EA Play e Minecraft. Os valores usam a base completa."),
        (26, "Contexto", "Janeiro e fevereiro têm apenas 2 registros cada; dezembro termina em 16/12. Não interprete a evolução mensal como série completa de demanda."),
        (29, "Compatibilidade", "Recomendado Microsoft Excel 2016 ou superior. Não há macros, APIs, conexões externas ou necessidade de Atualizar Tudo."),
        (32, "Reprodução", "Execute scripts/gerar_dashboard.py, scripts/gerar_previas.py e scripts/verificar_dashboard.py após instalar requirements.txt."),
    ]
    for row, title, body in sections:
        merge_box(ws, f"B{row}:B{row+1}", title, bg=SLATE, color=WHITE, size=9, bold=True,
                  align="left")
        merge_box(ws, f"C{row}:C{row+1}", body, bg=WHITE, color=TEXT, size=9, align="left")
        ws.row_dimensions[row].height = 28
        ws.row_dimensions[row + 1].height = 28

    merge_box(ws, "B35:C37",
              f"Autoria: {AUTHOR}. Projeto educacional e não oficial. Xbox, Game Pass, EA Play e Minecraft são marcas de seus respectivos titulares.",
              bg=INK, color=WHITE, size=9, bold=True)
    merge_box(ws, "B39:C40", "VOLTAR AO DASHBOARD", bg=GREEN, color=INK, size=10, bold=True,
              hyperlink="#'Dashboard'!A1")
    ws.print_area = "A1:C40"


def build_support(wb: Workbook) -> None:
    ws = wb.create_sheet("Apoio")
    ws.sheet_view.showGridLines = False
    lists = {
        "lista_planos": ["Todos", "Core", "Standard", "Ultimate"],
        "lista_tipos": ["Todos", "Monthly", "Quarterly", "Annual"],
        "lista_renovacao": ["Todos", "Yes", "No"],
    }
    ws["A1"] = "Listas internas para filtros"
    ws["A1"].fill = fill(INK)
    ws["A1"].font = Font(name="Aptos", size=12, bold=True, color=WHITE)
    for col, (name, values) in enumerate(lists.items(), 1):
        ws.cell(3, col, name)
        ws.cell(3, col).fill = fill(SLATE)
        ws.cell(3, col).font = Font(name="Aptos", size=8, bold=True, color=WHITE)
        for row, value in enumerate(values, 4):
            ws.cell(row, col, value)
        letter = chr(64 + col)
        add_name(wb, name, f"'Apoio'!${letter}$4:${letter}${3 + len(values)}")
        ws.column_dimensions[letter].width = 20
    ws.sheet_state = "hidden"


def add_defined_names(wb: Workbook) -> None:
    add_name(wb, "filtro_plano", "'Dashboard'!$C$6")
    add_name(wb, "filtro_tipo", "'Dashboard'!$G$6")
    add_name(wb, "filtro_renovacao", "'Dashboard'!$K$6")
    ranges = {
        "base_id": "'Base'!$A$5:$A$299",
        "base_plano": "'Base'!$C$5:$C$299",
        "base_inicio": "'Base'!$D$5:$D$299",
        "base_renovacao": "'Base'!$E$5:$E$299",
        "base_preco": "'Base'!$F$5:$F$299",
        "base_tipo": "'Base'!$G$5:$G$299",
        "base_ea": "'Base'!$I$5:$I$299",
        "base_minecraft": "'Base'!$K$5:$K$299",
        "base_cupom": "'Base'!$L$5:$L$299",
        "base_total": "'Base'!$M$5:$M$299",
    }
    for name, reference in ranges.items():
        add_name(wb, name, reference)


def build_workbook(records: list[list], summary: dict) -> Workbook:
    wb = Workbook()
    wb.properties.title = "Xbox Sales Pulse — Dashboard de Vendas em Excel"
    wb.properties.subject = "Projeto autoral para o desafio DIO de dashboard de vendas do Xbox"
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR_HANDLE
    wb.properties.description = "Dashboard de assinaturas com KPIs, filtros, gráficos e auditoria de dados."
    wb.properties.keywords = "Excel, Xbox, Game Pass, vendas, dashboard, DIO"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    # Cria primeiro a camada de análise, pois os gráficos do Dashboard
    # precisam referenciá-la durante a montagem. A aba ativa continua sendo
    # a primeira e recebe o Dashboard em seguida.
    wb.active.title = "Dashboard"
    build_analysis(wb)
    build_dashboard(wb, summary)
    build_base(wb, records)
    build_dictionary(wb)
    build_guide(wb)
    build_support(wb)
    add_defined_names(wb)
    wb.active = 0
    return wb


def cached_values(records: list[list], summary: dict) -> dict[str, dict[str, float | int | str]]:
    dashboard: dict[str, float | int | str] = {
        "A4": f"Base didática: {summary['count']} assinaturas • período 01/01/2024 a 16/12/2024 • valores em US$",
        "A8": "Filtros ativos: Plano = Todos • Periodicidade = Todos • Renovação = Todos",
        "A11": summary["revenue"],
        "D11": summary["count"],
        "G11": summary["ticket"],
        "J11": summary["auto_rate"],
        "M11": summary["validation_ok"] / summary["count"],
        "A16": summary["base"],
        "D16": summary["ea"],
        "G16": summary["minecraft"],
        "J16": summary["coupons"],
        "M16": summary["count"],
    }

    analysis: dict[str, float | int | str] = {
        "A4": "Filtros: Todos | Todos | Todos",
        "B6": summary["revenue"],
        "B7": summary["count"],
        "B8": summary["ticket"],
        "B9": summary["auto_rate"],
        "B10": summary["base"],
        "B11": summary["ea"],
        "B12": summary["minecraft"],
        "B13": summary["coupons"],
        "B14": summary["validation_ok"] / summary["count"],
        "B17": summary["annual_total"],
        "B18": summary["annual_yes"],
        "B19": summary["annual_no"],
        "B20": summary["ea"],
        "B21": summary["minecraft"],
    }
    for month in range(1, 13):
        row = 25 + month
        analysis[f"B{row}"] = summary["monthly"][month]["revenue"]
        analysis[f"C{row}"] = summary["monthly"][month]["count"]
    for row, plan in enumerate(("Core", "Standard", "Ultimate"), 26):
        analysis[f"F{row}"] = summary["plans"][plan]["revenue"]
        analysis[f"G{row}"] = summary["plans"][plan]["count"]
    for row, kind in enumerate(("Monthly", "Quarterly", "Annual"), 33):
        analysis[f"F{row}"] = summary["types"][kind]["revenue"]
        analysis[f"G{row}"] = summary["types"][kind]["count"]
    for row, status in enumerate(("Yes", "No"), 26):
        analysis[f"J{row}"] = summary["renewal"][status]["revenue"]
        analysis[f"K{row}"] = summary["renewal"][status]["count"]
    for row, plan in enumerate(("Core", "Standard", "Ultimate"), 33):
        analysis[f"J{row}"] = summary["plans"][plan]["ea"]
        analysis[f"K{row}"] = summary["plans"][plan]["minecraft"]
        analysis[f"L{row}"] = summary["plans"][plan]["revenue"]

    base: dict[str, float | int | str] = {}
    for row_index, record in enumerate(records, 5):
        calculated = record[5] + record[8] + record[10] - record[11]
        base[f"N{row_index}"] = calculated
        base[f"O{row_index}"] = "OK" if abs(calculated - record[12]) < 1e-9 else "REVIEW"
        month_start = date(record[3].year, record[3].month, 1)
        base[f"P{row_index}"] = to_excel(month_start)

    return {
        "xl/worksheets/sheet1.xml": dashboard,
        "xl/worksheets/sheet2.xml": analysis,
        "xl/worksheets/sheet3.xml": base,
    }


def populate_formula_cache(path: Path, records: list[list], summary: dict) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)
    sheets = cached_values(records, summary)
    temporary = path.with_name(path.stem + "-cache.xlsx")

    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in sheets:
                root = ET.fromstring(data)
                values = sheets[item.filename]
                for cell in root.findall(f".//{{{namespace}}}c"):
                    reference = cell.get("r")
                    if reference not in values:
                        continue
                    value = values[reference]
                    node = cell.find(f"{{{namespace}}}v")
                    if node is None:
                        node = ET.SubElement(cell, f"{{{namespace}}}v")
                    if isinstance(value, str):
                        cell.set("t", "str")
                        node.text = value
                    else:
                        if cell.get("t") == "str":
                            cell.attrib.pop("t", None)
                        node.text = format(value, ".15g")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    temporary.replace(path)


def main() -> None:
    records = load_records()
    write_normalized_csv(records)
    summary = summarize(records)
    wb = build_workbook(records, summary)
    wb.save(OUTPUT)
    populate_formula_cache(OUTPUT, records, summary)
    print(f"Dashboard criado: {OUTPUT}")
    print(f"CSV normalizado: {CSV_OUTPUT}")
    print(f"Registros: {summary['count']} | Receita: $ {summary['revenue']:,.2f}")


if __name__ == "__main__":
    main()
