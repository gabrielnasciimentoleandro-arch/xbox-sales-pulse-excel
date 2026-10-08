#!/usr/bin/env python3
"""Gera prévias autorais do Xbox Sales Pulse para o README."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyBboxPatch, Rectangle
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILE = ROOT / "xbox_sales_pulse.xlsx"
DOCS = ROOT / "docs"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"

INK = "#0B0F14"
DARK = "#111827"
SLATE = "#1F2937"
GREEN = "#22C55E"
LIME = "#9BC848"
MINT = "#2AE6B1"
BLUE = "#38BDF8"
PURPLE = "#8B5CF6"
AMBER = "#F59E0B"
RED = "#EF4444"
WHITE = "#FFFFFF"
TEXT = "#172033"
MUTED = "#64748B"
LIGHT = "#F4F7FA"
BORDER = "#CBD5E1"
PALE_GREEN = "#DCFCE7"


def money(value: float, decimals: bool = False) -> str:
    places = 2 if decimals else 0
    raw = f"{value:,.{places}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"$ {raw}"


def pct(value: float) -> str:
    return f"{value * 100:.1f}%".replace(".", ",")


def box(ax, x, y, w, h, face=WHITE, edge=BORDER, radius=.012, lw=1):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.004,rounding_size={radius}",
        facecolor=face, edgecolor=edge, linewidth=lw,
        transform=ax.transAxes, clip_on=False,
    )
    ax.add_patch(patch)
    return patch


def preview_dashboard(wb) -> None:
    dashboard = wb["Dashboard"]
    analysis = wb["Análise"]
    revenue = dashboard["A11"].value
    subscribers = dashboard["D11"].value
    ticket = dashboard["G11"].value
    auto_rate = dashboard["J11"].value
    valid_rate = dashboard["M11"].value
    base_revenue = dashboard["A16"].value
    ea = dashboard["D16"].value
    minecraft = dashboard["G16"].value
    coupons = dashboard["J16"].value

    months = [analysis.cell(row, 1).value.strftime("%b") for row in range(26, 38)]
    monthly_revenue = [analysis.cell(row, 2).value for row in range(26, 38)]
    plans = [analysis.cell(row, 5).value for row in range(26, 29)]
    plan_revenue = [analysis.cell(row, 6).value for row in range(26, 29)]
    types = [analysis.cell(row, 5).value for row in range(33, 36)]
    type_revenue = [analysis.cell(row, 6).value for row in range(33, 36)]
    addon_plans = [analysis.cell(row, 9).value for row in range(33, 36)]
    addon_ea = [analysis.cell(row, 10).value for row in range(33, 36)]
    addon_mc = [analysis.cell(row, 11).value for row in range(33, 36)]

    fig = plt.figure(figsize=(16, 10), facecolor="#E7EDF3")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_axis_off()
    box(ax, .018, .018, .964, .964, WHITE, "#B9C6D4", .01)

    ax.add_patch(Rectangle((.018, .875), .964, .107, transform=ax.transAxes,
                           facecolor=INK, edgecolor=INK))
    ax.text(.048, .935, "XBOX SALES PULSE", transform=ax.transAxes, color=WHITE,
            fontsize=26, fontweight="bold", va="center")
    ax.text(.048, .895, "Assinaturas • receita • recorrência • adicionais",
            transform=ax.transAxes, color="#B8C7D9", fontsize=10.5, va="center")
    ax.text(.95, .935, "PROJETO AUTORAL", transform=ax.transAxes, color=INK,
            fontsize=8.5, fontweight="bold", ha="right", va="center",
            bbox=dict(boxstyle="round,pad=.5", facecolor=GREEN, edgecolor=GREEN))
    ax.text(.95, .895, "295 registros • 2024 • US$", transform=ax.transAxes,
            color=WHITE, fontsize=9, ha="right", va="center")

    ax.text(.048, .842, "FILTROS", transform=ax.transAxes, color=TEXT,
            fontsize=9.5, fontweight="bold")
    filters = [("PLANO", "Todos"), ("PERIODICIDADE", "Todos"), ("RENOVAÇÃO", "Todos")]
    x = .12
    for label, value in filters:
        box(ax, x, .815, .195, .05, PALE_GREEN, "#8FE3AA")
        ax.text(x + .012, .84, label, transform=ax.transAxes, color=MUTED,
                fontsize=7.8, fontweight="bold", va="center")
        ax.text(x + .18, .84, value, transform=ax.transAxes, color=TEXT,
                fontsize=9.5, fontweight="bold", ha="right", va="center")
        x += .225

    cards = [
        ("RECEITA TOTAL", money(revenue), GREEN),
        ("ASSINANTES", f"{subscribers}", BLUE),
        ("TICKET MÉDIO", money(ticket, True), MINT),
        ("AUTO RENOVAÇÃO", pct(auto_rate), LIME),
        ("BASE VALIDADA", pct(valid_rate), PURPLE),
    ]
    card_x = [.048, .235, .422, .609, .796]
    card_w = .165
    for x, (label, value, accent) in zip(card_x, cards):
        box(ax, x, .69, card_w, .1, WHITE)
        ax.add_patch(Rectangle((x, .765), card_w, .025, transform=ax.transAxes,
                               facecolor=SLATE, edgecolor=SLATE))
        ax.text(x + card_w / 2, .7775, label, transform=ax.transAxes, color=WHITE,
                fontsize=7.3, fontweight="bold", ha="center", va="center")
        ax.text(x + card_w / 2, .727, value, transform=ax.transAxes, color=accent,
                fontsize=15.5, fontweight="bold", ha="center", va="center")

    component_cards = [
        ("Planos", money(base_revenue), GREEN),
        ("EA Play", money(ea), PURPLE),
        ("Minecraft", money(minecraft), MINT),
        ("Cupons", money(coupons), AMBER),
    ]
    x = .048
    for label, value, color in component_cards:
        box(ax, x, .625, .21, .045, LIGHT)
        ax.text(x + .012, .648, label, transform=ax.transAxes, color=MUTED,
                fontsize=8.2, va="center")
        ax.text(x + .195, .648, value, transform=ax.transAxes, color=color,
                fontsize=9.5, fontweight="bold", ha="right", va="center")
        x += .225

    line_ax = fig.add_axes([.055, .345, .43, .245], facecolor=WHITE)
    line_ax.plot(months, monthly_revenue, color=GREEN, marker="o", linewidth=2.8, markersize=5)
    line_ax.fill_between(months, monthly_revenue, color=GREEN, alpha=.1)
    line_ax.set_title("Receita por mês", loc="left", color=TEXT, fontsize=10.5,
                      fontweight="bold", pad=8)
    line_ax.grid(axis="y", color="#E5EAF0", linewidth=.8)
    line_ax.spines[["top", "right", "left"]].set_visible(False)
    line_ax.spines["bottom"].set_color(BORDER)
    line_ax.tick_params(colors=MUTED, labelsize=7.5)
    line_ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:,.0f}"))

    plan_ax = fig.add_axes([.56, .345, .39, .245], facecolor=WHITE)
    plan_colors = [BLUE, MINT, GREEN]
    y_pos = np.arange(len(plans))
    plan_ax.barh(y_pos, plan_revenue, color=plan_colors, height=.5)
    plan_ax.set_yticks(y_pos, plans)
    plan_ax.invert_yaxis()
    plan_ax.set_title("Receita por plano", loc="left", color=TEXT, fontsize=10.5,
                      fontweight="bold", pad=8)
    plan_ax.grid(axis="x", color="#E5EAF0", linewidth=.8)
    plan_ax.spines[["top", "right", "left"]].set_visible(False)
    plan_ax.spines["bottom"].set_color(BORDER)
    plan_ax.tick_params(colors=MUTED, labelsize=8)
    plan_ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:,.0f}"))
    for idx, value in enumerate(plan_revenue):
        plan_ax.text(value + 50, idx, money(value), color=TEXT, va="center",
                     fontsize=8, fontweight="bold")

    type_ax = fig.add_axes([.055, .075, .43, .21], facecolor=WHITE)
    type_colors = [BLUE, PURPLE, GREEN]
    bars = type_ax.bar(types, type_revenue, color=type_colors, width=.55)
    type_ax.set_title("Receita por periodicidade", loc="left", color=TEXT,
                      fontsize=10.5, fontweight="bold", pad=8)
    type_ax.grid(axis="y", color="#E5EAF0", linewidth=.8)
    type_ax.spines[["top", "right", "left"]].set_visible(False)
    type_ax.spines["bottom"].set_color(BORDER)
    type_ax.tick_params(colors=MUTED, labelsize=8)
    type_ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:,.0f}"))
    for bar, value in zip(bars, type_revenue):
        type_ax.text(bar.get_x() + bar.get_width() / 2, value + 45, money(value),
                     ha="center", color=TEXT, fontsize=7.5, fontweight="bold")

    addon_ax = fig.add_axes([.56, .075, .39, .21], facecolor=WHITE)
    idx = np.arange(len(addon_plans))
    width = .34
    addon_ax.bar(idx - width / 2, addon_ea, width, label="EA Play", color=PURPLE)
    addon_ax.bar(idx + width / 2, addon_mc, width, label="Minecraft", color=MINT)
    addon_ax.set_xticks(idx, addon_plans)
    addon_ax.set_title("Adicionais por plano", loc="left", color=TEXT,
                       fontsize=10.5, fontweight="bold", pad=8)
    addon_ax.grid(axis="y", color="#E5EAF0", linewidth=.8)
    addon_ax.spines[["top", "right", "left"]].set_visible(False)
    addon_ax.spines["bottom"].set_color(BORDER)
    addon_ax.tick_params(colors=MUTED, labelsize=8)
    addon_ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:,.0f}"))
    addon_ax.legend(frameon=False, fontsize=8, loc="upper left")

    ax.text(.048, .035,
            f"Perguntas-chave: Anual $1.754 • Anual c/ renovação $1.537 • Anual s/ renovação $217 • EA $2.940 • Minecraft $3.880",
            transform=ax.transAxes, color=TEXT, fontsize=8.6, fontweight="bold")
    ax.text(.95, .035, AUTHOR, transform=ax.transAxes, color=MUTED,
            fontsize=7.8, ha="right")

    output = DOCS / "dashboard-geral.png"
    fig.savefig(output, dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"Prévia criada: {output}")


def preview_quality(wb) -> None:
    analysis = wb["Análise"]
    base = wb["Base"]
    rows = []
    for row in range(5, 11):
        rows.append([base.cell(row, col).value for col in range(1, 14)])

    fig = plt.figure(figsize=(16, 9), facecolor="#E7EDF3")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_axis_off()
    box(ax, .02, .025, .96, .95, WHITE, "#B9C6D4", .01)
    ax.add_patch(Rectangle((.02, .84), .96, .135, transform=ax.transAxes,
                           facecolor=INK, edgecolor=INK))
    ax.text(.05, .915, "XBOX SALES PULSE  /  DATA VIEW", transform=ax.transAxes,
            color=WHITE, fontsize=23, fontweight="bold", va="center")
    ax.text(.05, .87, "Qualidade da base e respostas das perguntas de negócio",
            transform=ax.transAxes, color="#B8C7D9", fontsize=10.5, va="center")
    ax.text(.95, .915, "295 / 295", transform=ax.transAxes, color=GREEN,
            fontsize=21, fontweight="bold", ha="right", va="center")
    ax.text(.95, .875, "valores conferidos", transform=ax.transAxes, color=WHITE,
            fontsize=8.5, ha="right")

    kpis = [
        ("Receita total", money(analysis["B6"].value), GREEN),
        ("Plano anual", money(analysis["B17"].value), BLUE),
        ("Anual + renovação", money(analysis["B18"].value), LIME),
        ("Anual sem renovação", money(analysis["B19"].value), RED),
        ("EA Play", money(analysis["B20"].value), PURPLE),
        ("Minecraft", money(analysis["B21"].value), MINT),
    ]
    x_positions = [.05, .205, .36, .515, .67, .825]
    for x, (label, value, color) in zip(x_positions, kpis):
        box(ax, x, .71, .135, .09, LIGHT)
        ax.text(x + .0675, .773, label, transform=ax.transAxes, color=MUTED,
                fontsize=7.5, ha="center", va="center")
        ax.text(x + .0675, .738, value, transform=ax.transAxes, color=color,
                fontsize=11.5, fontweight="bold", ha="center", va="center")

    headers = ["ID", "Nome", "Plano", "Início", "Renovação", "Tipo", "Total"]
    col_x = [.05, .12, .30, .42, .53, .65, .82]
    col_w = [.06, .17, .11, .10, .11, .16, .12]
    header_y = .64
    for x, w, label in zip(col_x, col_w, headers):
        ax.add_patch(Rectangle((x, header_y), w, .045, transform=ax.transAxes,
                               facecolor=SLATE, edgecolor=WHITE, linewidth=.5))
        ax.text(x + .006, header_y + .0225, label, transform=ax.transAxes,
                color=WHITE, fontsize=8, fontweight="bold", va="center")

    y = .56
    row_h = .075
    for idx, row in enumerate(rows):
        face = WHITE if idx % 2 == 0 else LIGHT
        ax.add_patch(Rectangle((.05, y), .89, row_h - .006, transform=ax.transAxes,
                               facecolor=face, edgecolor=BORDER, linewidth=.6))
        values = [
            row[0], row[1], row[2], row[3].strftime("%d/%m/%Y"), row[4], row[6], money(row[12]),
        ]
        for x, value in zip(col_x, values):
            text = str(value)
            if len(text) > 24:
                text = text[:22] + "…"
            ax.text(x + .006, y + (row_h - .006) / 2, text, transform=ax.transAxes,
                    color=TEXT, fontsize=8.2, va="center")
        y -= row_h

    box(ax, .05, .075, .89, .075, PALE_GREEN, "#8FE3AA")
    ax.text(.07, .122, "AUDITORIA", transform=ax.transAxes, color=TEXT,
            fontsize=9.5, fontweight="bold")
    ax.text(.07, .094,
            "Total Value = Subscription Price + EA Play + Minecraft − Coupon em todos os 295 registros. IDs são únicos; dezembro é parcial até 16/12/2024.",
            transform=ax.transAxes, color=TEXT, fontsize=8.5)
    ax.text(.05, .045,
            f"Base didática DIO • versão normalizada e painel autoral • {AUTHOR}",
            transform=ax.transAxes, color=MUTED, fontsize=8)

    output = DOCS / "data-view.png"
    fig.savefig(output, dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"Prévia criada: {output}")


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    if not FILE.exists():
        raise SystemExit("Gere o dashboard antes de criar as prévias.")
    wb = load_workbook(FILE, data_only=True)
    preview_dashboard(wb)
    preview_quality(wb)


if __name__ == "__main__":
    main()
